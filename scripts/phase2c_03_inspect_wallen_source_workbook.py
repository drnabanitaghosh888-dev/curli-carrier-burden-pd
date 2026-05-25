#!/usr/bin/env python3
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook


def infer_type(name: str) -> str:
    n = name.lower()
    if "sample" in n and "meta" in n:
        return "sample_metadata_candidate"
    if "species" in n:
        return "species_abundance_candidate"
    if "genus" in n:
        return "genus_abundance_candidate"
    if "gene" in n or "family" in n:
        return "gene_family_candidate"
    if "pathway" in n:
        return "pathway_candidate"
    if "clinical" in n or "phenotype" in n:
        return "clinical_metadata_candidate"
    return "unknown"


def row_preview(ws, n: int = 20) -> str:
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        vals = [str(x) if x is not None else "" for x in row[:n]]
        return "|".join(vals)
    return ""


def first_nonempty_row_index(ws, max_scan: int = 2000) -> int:
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan, values_only=True), start=1):
        if any(cell not in (None, "") for cell in row):
            return i
    return -1


def main() -> int:
    root = Path(__file__).resolve().parents[1]
    workbook = root / "data/phase2c/manual_sources/wallen_prjna834801/Source_Data_24Oct2022.xlsx"
    out_dir = root / "data/phase2c/reports"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_tsv = out_dir / "phase2c_wallen_workbook_sheet_inventory.tsv"
    out_status = out_dir / "phase2c_wallen_workbook_status.txt"

    if not workbook.exists():
        with out_tsv.open("w", encoding="utf-8", newline="") as h:
            w = csv.DictWriter(
                h,
                fieldnames=[
                    "sheet_name",
                    "max_row",
                    "max_column",
                    "first_nonempty_row_index",
                    "first_row_preview",
                    "candidate_table_type",
                ],
                delimiter="\t",
            )
            w.writeheader()
        out_status.write_text("MISSING_WORKBOOK\n", encoding="utf-8")
        print(f"[DONE] wrote: {out_tsv}")
        print(f"[DONE] wrote: {out_status}")
        return 0

    wb = load_workbook(workbook, read_only=True, data_only=True)
    rows = []
    for s in wb.sheetnames:
        ws = wb[s]
        rows.append(
            {
                "sheet_name": s,
                "max_row": str(ws.max_row or 0),
                "max_column": str(ws.max_column or 0),
                "first_nonempty_row_index": str(first_nonempty_row_index(ws)),
                "first_row_preview": row_preview(ws, n=20),
                "candidate_table_type": infer_type(s),
            }
        )

    with out_tsv.open("w", encoding="utf-8", newline="") as h:
        w = csv.DictWriter(
            h,
            fieldnames=[
                "sheet_name",
                "max_row",
                "max_column",
                "first_nonempty_row_index",
                "first_row_preview",
                "candidate_table_type",
            ],
            delimiter="\t",
        )
        w.writeheader()
        for r in rows:
            w.writerow(r)

    out_status.write_text("WORKBOOK_INVENTORY_READY\n", encoding="utf-8")
    print(f"[DONE] wrote: {out_tsv}")
    print(f"[DONE] wrote: {out_status}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
